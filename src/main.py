import csv
from pathlib import Path

from src.generators import filter_by_currency
from src.processing import filter_by_state, sort_by_date
from src.services import process_bank_search
from src.utils import load_operations
from src.widget import get_date, mask_account_card

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"

JSON_PATH = DATA_DIR / "operations.json"
CSV_PATH = DATA_DIR / "transactions.csv"
XLSX_PATH = DATA_DIR / "transactions_excel.xlsx"


def read_csv_file(file_path: str | Path) -> list[dict]:
    """Читает транзакции из CSV-файла."""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return list(csv.DictReader(f, delimiter=";"))
    except FileNotFoundError:
        print(f"Файл не найден: {file_path}")
        return []


def read_xlsx_file(file_path: str | Path) -> list[dict]:
    """Читает транзакции из XLSX-файла."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("Установите openpyxl: pip install openpyxl")
        return []

    try:
        wb = load_workbook(file_path)
    except FileNotFoundError:
        print(f"Файл не найден: {file_path}")
        return []

    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]


def ask_yes_no(prompt: str) -> bool:
    """Запрашивает у пользователя ответ Да/Нет."""
    while True:
        answer = input(prompt).strip().lower()
        if answer in ("да", "yes", "y", "д"):
            return True
        if answer in ("нет", "no", "n", "н"):
            return False
        print('Пожалуйста, введите "Да" или "Нет".')


def _get_amount_and_currency(transaction: dict) -> tuple:
    """Достаёт сумму и код валюты из транзакции."""
    amount = transaction.get("amount")
    currency = transaction.get("currency")
    if amount is not None and currency is not None:
        return amount, currency

    op_amount = transaction.get("operationAmount", {})
    if isinstance(op_amount, dict):
        amount = op_amount.get("amount", "")
        currency_obj = op_amount.get("currency", {})
        currency = currency_obj.get("code", "") if isinstance(currency_obj, dict) else ""
        return amount, currency

    return "", ""


def print_transaction(transaction: dict) -> None:
    """Печатает одну транзакцию в читаемом виде."""
    date = get_date(transaction.get("date", ""))
    description = transaction.get("description", "")
    from_acc = transaction.get("from", "")
    to_acc = transaction.get("to", "")
    amount, currency = _get_amount_and_currency(transaction)

    print(f"{date} {description}")

    if from_acc and to_acc:
        print(f"{mask_account_card(from_acc)} -> " f"{mask_account_card(to_acc)}")
    elif to_acc:
        print(mask_account_card(to_acc))

    currency_suffix = "руб." if currency == "RUB" else currency
    print(f"Сумма: {amount} {currency_suffix}\n")


def _load_transactions(choice: str) -> list[dict]:
    """Загружает транзакции из файла по выбору пользователя."""
    if choice == "1":
        print("Для обработки выбран JSON-файл.")
        return load_operations(JSON_PATH)

    if choice == "2":
        print("Для обработки выбран CSV-файл.")
        return read_csv_file(CSV_PATH)

    if choice == "3":
        print("Для обработки выбран XLSX-файл.")
        return read_xlsx_file(XLSX_PATH)

    return []


def _choose_transactions() -> list[dict]:
    """Запрашивает источник данных и загружает транзакции."""
    while True:
        choice = input().strip()
        if choice in ("1", "2", "3"):
            return _load_transactions(choice)
        print("Некорректный выбор. Введите 1, 2 или 3.")


def _choose_status() -> str:
    """Запрашивает статус фильтрации."""
    available = ["EXECUTED", "CANCELED", "PENDING"]
    while True:
        print("Введите статус, по которому необходимо выполнить " "фильтрацию.")
        print(f"Доступные для фильтровки статусы: " f'{", ".join(available)}')

        status = input().strip().upper()
        if status in available:
            print(f'Операции отфильтрованы по статусу "{status}"')
            return status

        print(f'Статус операции "{status}" недоступен.')


def _apply_sort(transactions: list[dict]) -> list[dict]:
    """Спрашивает про сортировку и применяет её."""
    if not ask_yes_no("Отсортировать операции по дате? Да/Нет\n"):
        return transactions

    order = input("Отсортировать по возрастанию или по убыванию?\n").strip().lower()

    if order in ("по возрастанию", "возрастание", "asc"):
        return sort_by_date(transactions, reverse=False)
    if order in ("по убыванию", "убывание", "desc"):
        return sort_by_date(transactions, reverse=True)

    print("Некорректный ввод, сортировка не применена.")
    return transactions


def _apply_filters(transactions: list[dict]) -> list[dict]:
    """Применяет фильтр по RUB и по слову в описании."""
    if ask_yes_no("Выводить только рублевые транзакции? Да/Нет\n"):
        transactions = list(filter_by_currency(transactions, "RUB"))

    if ask_yes_no("Отфильтровать список транзакций по определенному слову " "в описании? Да/Нет\n"):
        search = input("Введите слово для поиска: ").strip()
        transactions = process_bank_search(transactions, search)

    return transactions


def main() -> None:
    """Основная функция проекта: связывает всю логику вместе."""
    print("Привет! Добро пожаловать в программу работы " "с банковскими транзакциями.")
    print("Выберите необходимый пункт меню:")
    print("1. Получить информацию о транзакциях из JSON-файла")
    print("2. Получить информацию о транзакциях из CSV-файла")
    print("3. Получить информацию о транзакциях из XLSX-файла")

    transactions = _choose_transactions()
    if not transactions:
        print("Не удалось загрузить транзакции.")
        return

    status = _choose_status()
    transactions = filter_by_state(transactions, status)

    transactions = _apply_sort(transactions)
    transactions = _apply_filters(transactions)

    print("Распечатываю итоговый список транзакций...")

    if not transactions:
        print("Не найдено ни одной транзакции, подходящей под ваши " "условия фильтрации")
        return

    print(f"\nВсего банковских операций в выборке: " f"{len(transactions)}\n")

    for transaction in transactions:
        print_transaction(transaction)


if __name__ == "__main__":
    main()
